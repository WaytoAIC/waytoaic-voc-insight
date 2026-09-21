#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Reddit 线（进阶版）· 第 3 步 算数，出报告

Reddit 上什么人都有，说的也不一定是你的货。所以报告的主线是：**先分流，再各看各的**——
还没买的人想解决什么、看重什么；买了的人（我的货 / 对手的货）好在哪不好在哪；放弃的人为什么放弃、换成了什么。

点赞数只当"有多少人有同感"，不当"有多少人会买"。
分工：脚本不编故事，AI 不算数字。
"""
import argparse, html as H, json, os, sys
from voc_common import (pd, load_config, multi, split_cell, base_tag, tag_dir, all_tags, dir_defs, DIR_SEP, dump)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

NEED = ["序号", "群组", "原话", "帖子类型", "到哪一步", "说谁的货", "想解决什么", "买之前看重什么", "现在怎么凑合",
        "好评的事", "差评的事", "客户的说法", "该谁改", "多严重", "整体好坏"]
E = lambda s: H.escape(str(s if s is not None else ""))
F = lambda **k: Font(name="Arial", **k)


def validate(d, cfg, sku):
    err, rc = [], cfg["Reddit"]
    miss = [c for c in NEED if c not in d.columns]
    if miss:
        sys.exit(f"[打标表缺格子] {miss}")
    for colname, allowed in (("帖子类型", rc["帖子类型"]), ("到哪一步", rc["到哪一步"]), ("说谁的货", rc["说谁的货"])):
        bad = sorted(set(d[colname]) - set(allowed))
        if bad:
            err.append(f"「{colname}」只能填 {allowed}，出现了 {bad}")
    for colname, allowed in (("想解决什么", rc["想解决什么"]), ("现在怎么凑合", rc["现在怎么凑合"])):
        bad = sorted({t for v in d[colname] for t in split_cell(v)} - set(allowed))
        if bad:
            err.append(f"「{colname}」里有配置没定义的：{bad}（确实是新的，先记成下一版的提案）")
    ok, dirs = all_tags(cfg, (sku or {}).get("这一款特有的标签")), dir_defs(cfg)
    unknown, bad_dir = set(), set()
    for colname in ("买之前看重什么", "好评的事", "差评的事"):
        for v in d[colname]:
            for t in split_cell(v):
                b, dr = base_tag(t), tag_dir(t)
                if b not in ok:
                    unknown.add(b)
                elif dr and (colname != "差评的事" or b not in dirs or dr not in dirs[b]):
                    bad_dir.add(f"{colname}:{t}")
    if unknown:
        err.append("这些细项不在标签表里：" + "、".join(sorted(unknown)))
    if bad_dir:
        err.append("方向写法不对：" + "、".join(sorted(bad_dir)))
    mine_neg = (d["说谁的货"] == "我的") & (d["差评的事"] != "")
    attrs = [a["名称"] for a in cfg["该谁改"]]
    if (mine_neg & ~d["该谁改"].isin(attrs)).any():
        err.append("说的是我的货、且有差评的事的行，「该谁改」必须填")
    if ((~mine_neg) & (d["该谁改"] != "")).any():
        err.append("只有说的是我的货、且有差评的事，才填「该谁改」「多严重」；别的行留空（对手的毛病不归我们谁改）")
    if err:
        sys.exit("[打标表没过校验]\n- " + "\n- ".join(dict.fromkeys(err)))


def counts_with_votes(d, col):
    rows = {}
    for _, r in d.iterrows():
        for t in split_cell(r[col]):
            b = base_tag(t)
            x = rows.setdefault(b, {"n": 0, "votes": 0})
            x["n"] += 1
            x["votes"] += int(r.get("点赞", 0) or 0)
    return sorted(({"tag": k, **v} for k, v in rows.items()), key=lambda x: (-x["n"], -x["votes"]))


def pos_neg(d):
    p, n = multi(d["好评的事"]), multi(d["差评的事"])
    p = p.groupby(p.index.map(base_tag)).sum() if len(p) else p
    n = n.groupby(n.index.map(base_tag)).sum() if len(n) else n
    tags = sorted(set(p.index) | set(n.index), key=lambda t: -(p.get(t, 0) + n.get(t, 0)))
    return [{"tag": t, "好": int(p.get(t, 0)), "差": int(n.get(t, 0))} for t in tags]


def compute(d, prep, cfg):
    rc = cfg["Reddit"]
    d = d.copy()
    d["点赞"] = pd.to_numeric(d.get("点赞", 0), errors="coerce").fillna(0).astype(int)
    noise = d["帖子类型"] == "广告或无关"
    a = d[~noise]
    stages, whose = rc["到哪一步"], rc["说谁的货"]
    matrix = {s: {w: int(((a["到哪一步"] == s) & (a["说谁的货"] == w)).sum()) for w in whose} for s in stages}
    pre = a[a["到哪一步"].isin(["没买在问", "正在几家里挑"])]
    bought = a[a["到哪一步"].isin(["买了在用", "用过放弃"])]
    quit_ = a[a["到哪一步"] == "用过放弃"]
    mine, rival = bought[bought["说谁的货"] == "我的"], bought[bought["说谁的货"].isin(["对手的", "整个品类"])]
    heavy = {x["发言人"] for x in (prep or {}).get("发了 3 条以上的人", [])}
    top = a.sort_values("点赞", ascending=False).head(8)
    sym = []
    for nm, (x, y) in dir_defs(cfg).items():
        cx = int(a["差评的事"].map(lambda v: f"{nm}{DIR_SEP}{x}" in split_cell(v)).sum())
        cy = int(a["差评的事"].map(lambda v: f"{nm}{DIR_SEP}{y}" in split_cell(v)).sum())
        if cx + cy:
            sym.append({"细项": nm, x: cx, y: cy})
    phrases = multi(a["客户的说法"].str.lower())
    return {
        "n_tagged": len(d), "n": len(a), "noise": int(noise.sum()),
        "groups": a["群组"].value_counts().to_dict(),
        "post_types": a["帖子类型"].value_counts().to_dict(),
        "stage": a["到哪一步"].value_counts().to_dict(), "whose": a["说谁的货"].value_counts().to_dict(),
        "matrix": matrix, "stages": stages, "whose_order": whose,
        "n_pre": len(pre), "n_bought": len(bought), "n_quit": len(quit_), "n_mine": len(mine), "n_rival": len(rival),
        "need": counts_with_votes(a, "想解决什么"),
        "care": counts_with_votes(pre, "买之前看重什么"),
        "care_all": counts_with_votes(a, "买之前看重什么"),
        "workaround": counts_with_votes(a, "现在怎么凑合"),
        "mine": pos_neg(mine), "rival": pos_neg(rival),
        "quit_why": counts_with_votes(quit_, "差评的事"), "quit_to": counts_with_votes(quit_, "现在怎么凑合"),
        "n_mine_people": int(mine["发言人"].nunique()) if "发言人" in mine.columns else None,
        "n_rival_people": int(rival["发言人"].nunique()) if "发言人" in rival.columns else None,
        "n_people": int(a["发言人"].nunique()) if "发言人" in a.columns else None,
        "mine_attr": mine.loc[mine["差评的事"] != "", "该谁改"].value_counts().to_dict(),
        "mine_overall": mine["整体好坏"].value_counts().to_dict(),
        "rival_overall": rival["整体好坏"].value_counts().to_dict(),
        "phrases": [{"p": k, "n": int(v)} for k, v in phrases.items()],
        "top": [{"序号": int(r["序号"]), "点赞": int(r["点赞"]), "到哪一步": r["到哪一步"], "说谁的货": r["说谁的货"],
                 "原话": r["原话"], "反复发的人": r.get("发言人", "") in heavy} for _, r in top.iterrows()],
        "sym": sym, "prep": prep or {},
    }


CSS_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "review_report.css")
EXTRA_CSS = """
.mx{width:100%;border-collapse:collapse;margin-top:22px;font-size:14px}
.mx th{font-family:var(--mono);font-size:11px;letter-spacing:.08em;color:var(--mut);padding:8px 10px;text-align:center;border-bottom:1px solid var(--ink)}
.mx th:first-child,.mx td:first-child{text-align:left}
.mx td{padding:0;border-bottom:1px solid var(--line);text-align:center}
.mx td:first-child{padding:10px 10px 10px 0;font-weight:500}
.mx .cell{display:block;padding:10px 6px;font-family:var(--mono);font-size:13px}
.mx .sum{font-family:var(--mono);color:var(--mut)}
.two{display:grid;grid-template-columns:1fr 1fr;gap:26px;margin-top:8px}
@media (max-width:760px){.two{grid-template-columns:1fr}}
.two h3{margin-top:22px}
.vrow{display:grid;grid-template-columns:170px 1fr 96px;gap:12px;align-items:center;padding:6px 0;border-bottom:1px solid var(--line);font-size:13.5px}
.vrow .t{height:11px;background:#DDE2E2}.vrow .t i{display:block;height:100%;background:var(--ink)}
.vrow .v{font-family:var(--mono);font-size:12px;color:var(--mut);text-align:right}
.phr{display:flex;flex-wrap:wrap;gap:8px;margin-top:18px}
.phr span{font-family:var(--mono);font-size:12px;padding:4px 10px;border:1px solid var(--line);background:var(--card);border-radius:2px}
.phr span b{color:var(--wet);font-weight:600;margin-left:6px}
.topq{margin-top:20px}
.topq .up{font-family:var(--mono);font-size:12px;color:var(--wet);font-weight:600}
.rule{margin-top:26px;border:1px solid var(--ink2);padding:16px 20px;font-size:13.5px;color:var(--ink2);border-radius:2px}
"""


def vrows(items, key="n", label=lambda x: x["tag"], extra=lambda x: f'{x["n"]} 条 · {x["votes"]} 赞'):
    if not items:
        return '<p style="color:var(--mut);font-size:13.5px">这一轮没有。</p>'
    mx = max(x[key] for x in items) or 1
    return "".join(f'<div class="vrow"><div>{E(label(x))}</div><div class="t"><i style="width:{x[key]/mx*100:.0f}%"></i></div>'
                   f'<div class="v">{E(extra(x))}</div></div>' for x in items)


def pn_bars(rows):
    if not rows:
        return '<p style="color:var(--mut);font-size:13.5px">这一轮没有。</p>'
    out = ""
    for r in rows:
        tot = r["好"] + r["差"]
        out += (f'<div class="bar-row"><div class="bl">{E(r["tag"])}</div><div class="bt">'
                f'<span class="dry" style="--w:{r["好"]/tot*100:.1f}%"></span><span class="wet" style="--w:{r["差"]/tot*100:.1f}%"></span></div>'
                f'<div class="bn"><b>{tot}</b><i>好 {r["好"]} · 差 {r["差"]}</i></div></div>')
    return f'<div class="bars">{out}</div>'


def kpi5(t):
    if t.get("metric"):
        return (f'<b>看哪个数</b> {E(t["metric"])}<br><b>现在</b> {E(t.get("now","—"))} → <b>要到</b> {E(t.get("target","—"))}'
                f'<br><b>哪天看</b> {E(t.get("when","—"))}'
                + (f'<span class="ifnot"><b>要是没变</b> {E(t["if_not"])}</span>' if t.get("if_not") else ""))
    return E(t.get("kpi", ""))


def build_html(d, m, ins, path, product, cfg, is_sample):
    css = open(CSS_PATH, encoding="utf-8").read() + EXTRA_CSS
    prep = m["prep"]
    share = lambda k: f'{m["stage"].get(k,0)/max(1,m["n"])*100:.0f}%'
    meta = [("拿到的讨论", prep.get("原始条数", m["n_tagged"]), 0),
            ("剔掉删除、太短、广告跑题", f'−{prep.get("原始条数", m["n_tagged"]) - m["n"]}', 0),
            ("进入分析", m["n"], 0), ("几个群组", len(m["groups"]), 0),
            ("其中买了的人", f'{m["n_bought"]/max(1,m["n"])*100:.0f}%', 0),
            ("同一个人最多占", f'{(prep.get("同一个人最多占") or 0)*100:.0f}%', 1)]
    meta_h = "".join(f'<div class="{"neg" if b else ""}"><b>{v}</b><span>{k}</span></div>' for k, v, b in meta)
    mxv = max([v for row in m["matrix"].values() for v in row.values()] + [1])
    head = "".join(f"<th>{E(w)}</th>" for w in m["whose_order"]) + "<th>合计</th>"
    body = ""
    for s in m["stages"]:
        row = m["matrix"][s]
        if not sum(row.values()):
            continue
        cells = "".join(f'<td><span class="cell" style="background:rgba(34,118,107,{row[w]/mxv*0.55:.2f})">{row[w] or "·"}</span></td>'
                        for w in m["whose_order"])
        body += f'<tr><td>{E(s)}</td>{cells}<td class="sum">{sum(row.values())}</td></tr>'
    phr = "".join(f'<span>{E(x["p"])}<b>{x["n"]}</b></span>' for x in m["phrases"] if x["n"] >= 2) or \
          "".join(f'<span>{E(x["p"])}</span>' for x in m["phrases"][:24])
    phr_once = "".join(f'<span>{E(x["p"])}</span>' for x in m["phrases"] if x["n"] < 2)
    topq = "".join(f'<blockquote><p>{E(t["原话"])}</p><cite><span class="up">▲ {t["点赞"]}</span> · {E(t["到哪一步"])} · 说的是{E(t["说谁的货"])}'
                   f'{" · 这个人发了好几条，别当成好几个人" if t["反复发的人"] else ""}</cite></blockquote>' for t in m["top"][:6])
    lay = ins.get("layers", [])
    layers = "".join(f'<div class="layer {"last" if i==len(lay)-1 else ""}"><div class="q">{E(l.get("level",""))} · {E(l.get("q",""))}</div>'
                     f'<h4>{l["title"]}</h4><p>{l["body"]}</p></div>' for i, l in enumerate(lay))
    chains = "".join(f'<div class="chain"><div class="tag">{E(c.get("tag",""))}</div><h4>{c["title"]}</h4><p>{c["body"]}</p></div>'
                     for c in ins.get("chains", []))
    tiers = cfg.get("行动四档", ["马上止血", "根治问题", "防止再犯", "要老板定"])
    todo = ""
    for ti, tier in enumerate(tiers, 1):
        items = [t for t in ins.get("todo", []) if t.get("tier") == tier]
        if not items:
            continue
        cards = "".join(f'<article class="todo"><header><span class="own">{E(i.get("owner",""))} · 用在哪：{E(i.get("use",""))}'
                        f'{" · 依据 "+str(i["n"])+" 条" if i.get("n") else ""}</span><h4>{i.get("problem","")}</h4></header>'
                        f'<p>{i["body"]}</p><div class="kpi"><span>怎么才算做完</span>{kpi5(i)}</div></article>' for i in items)
        todo += (f'<div class="tier"><div class="tier-h"><span class="pill t{ti}">{E(tier)}</span>'
                 f'<span class="when">{E(ins.get("tier_when",{}).get(tier,""))}</span></div><div class="todo-grid">{cards}</div></div>')
    caveats = "".join(f"<li>{c}</li>" for c in ins.get("caveats", []))
    recs = d.drop(columns=[c for c in ("记录编号",) if c in d.columns]).fillna("").to_dict("records")
    payload = json.dumps(recs, ensure_ascii=False, default=str).replace("</", "<\\/")
    sec = lambda slab, h2, inner: f'<section><div class="slab">{slab}</div><h2>{h2}</h2>{inner}</section>'
    note = lambda k, default="": f'<p>{ins.get(k, default)}</p>' if ins.get(k, default) else ""
    page = f'''<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>{E(product)} · Reddit 讨论洞察报告</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+SC:wght@300;400;500;700&family=Noto+Serif+SC:wght@500;700;900&family=IBM+Plex+Mono:wght@400;500;600&display=swap" rel="stylesheet">
<style>{css}</style></head><body><div class="wrap">
<header class="top"><div class="eyebrow">Reddit 讨论洞察报告 · 进阶版{' <span class="sample-flag">示例数据</span>' if is_sample else ''}</div>
<h1>{ins.get("headline","")}</h1><p class="sub">{ins.get("subtitle","")}</p><div class="meta">{meta_h}</div></header>

<section><div class="slab">结论先行</div><div class="thesis"><p>{"<br>".join(ins.get("thesis", []))}</p></div>
<p class="lead" style="margin-top:26px">{ins.get("lead","")}</p></section>

{sec("先分流", "这些人到哪一步了，说的是谁的货", note("matrix_note") +
     f'<figure style="padding:8px 18px 14px"><table class="mx"><thead><tr><th></th>{head}</tr></thead><tbody>{body}</tbody></table>'
     f'<figcaption>颜色越深条数越多。退货和评论，平台已经替你确认了"这人买了、买的是你这一款"；Reddit 什么都没替你确认，所以要先分。</figcaption></figure>')}

{sec("还没买的人", f'{m["n_pre"]} 个还没买的人：想解决什么，看重什么',
     note("pre_note") + f'<div class="two"><div><h3>想解决什么</h3>{vrows(m["need"])}</div>'
     f'<div><h3>买之前看重什么</h3>{vrows(m["care"])}</div></div>')}

{sec("买了的人", f'{m["n_bought"]} 个买了的人：我的货，和对手的货', note("bought_note") +
     f'<div class="two"><div><h3>说我的货（{m["n_mine"]} 条 · {m.get("n_mine_people") or "?"} 个人）</h3>{pn_bars(m["mine"])}</div>'
     f'<div><h3>说对手的货和整个品类（{m["n_rival"]} 条 · {m.get("n_rival_people") or "?"} 个人）</h3>{pn_bars(m["rival"])}</div></div>'
     '<p style="margin-top:16px;font-size:13.5px;color:var(--mut)">深色是好评，砖红是差评。右边这一栏的砖红，就是对手的毛病。</p>')}

{sec("放弃的人 · 凑合的办法", f'{m["n_quit"]} 个放弃的人为什么放弃；大家现在拿什么凑合', note("quit_note") +
     f'<div class="two"><div><h3>为什么放弃</h3>{vrows(m["quit_why"])}</div>'
     f'<div><h3>现在拿什么凑合（全部人）</h3>{vrows(m["workaround"])}</div></div>')}

{sec("客户的说法", "他们用什么词说这件事", note("phrase_note") + f'<div class="phr">{phr}</div>' +
     (f'<p style="margin-top:18px;font-size:13px;color:var(--mut)">只出现过一次的：</p><div class="phr">{phr_once}</div>' if phr_once and any(x["n"]>=2 for x in m["phrases"]) else ""))}

{sec("同感最多的", "点赞最多的几条", note("top_note", "点赞数只当「有多少人有同感」，不当「有多少人会买」。") + f'<div class="topq">{topq}</div>')}

{sec("为什么", "几条因果链", f'<div class="chains">{chains}</div>') if chains else ""}
{sec("三层洞察", "发生了什么 · 为什么 · 不改会怎样", f'<div class="layers">{layers}</div>')}
{sec("行动清单", "下一步做什么", note("todo_note", "Reddit 这条线的出口跟退货、评论不一样：做什么品、主打什么、页面用什么词、内容选什么题；说我的货的差评，照样进「该谁改」。") + todo)}

<section><div class="slab">原话查阅</div><h2>{m["n_tagged"]} 条讨论，随时点回去看</h2>
<div class="filters" id="filters"></div><div class="vlist" id="vlist"></div></section>

<div class="caveat"><h3>这批数据答不了什么</h3><ul>{caveats}</ul></div>
<div class="rule"><b>怎么拿的、怎么用的：</b>人手动、低频、只读；作者名已去掉；不存进库、不设定时任务、不拿去训练模型。
平台条款禁止把这些讨论拿去做商业挖掘，规模化要走官方授权。</div>
<footer><span>{E(ins.get("sig") or cfg.get("图上文字",{}).get("落款",""))}</span><span>{m["n"]} 条 · {" / ".join(m["groups"])}</span></footer>
</div>
<script>
const RECS={payload};
const esc=s=>String(s==null?'':s).replace(/[<>&]/g,c=>({{'<':'&lt;','>':'&gt;','&':'&amp;'}}[c]));
const fb=document.getElementById('filters'),list=document.getElementById('vlist');
const opts=['全部'].concat({json.dumps(m["stages"], ensure_ascii=False)});let active='全部';
const pick=a=>a==='全部'?RECS:RECS.filter(r=>r['到哪一步']===a);
opts.forEach(a=>{{if(a!=='全部'&&!pick(a).length)return;const b=document.createElement('button');b.className='chip';b.type='button';
  b.textContent=`${{a}} ${{pick(a).length}}`;b.setAttribute('aria-pressed',a===active);
  b.onclick=()=>{{active=a;fb.querySelectorAll('.chip').forEach(c=>c.setAttribute('aria-pressed','false'));b.setAttribute('aria-pressed','true');render();}};
  fb.appendChild(b);}});
function render(){{list.innerHTML=pick(active).map(r=>`<div class="vcard"><span class="idx">#${{r['序号']}} · ▲${{esc(r['点赞'])}} · ${{esc(r['群组'])}} · ${{esc(r['到哪一步'])}} · 说的是${{esc(r['说谁的货'])}}</span>
  ${{esc(r['原话'])}}<div class="tags">${{r['想解决什么']?`想解决：${{esc(r['想解决什么'])}}<br>`:''}}${{r['买之前看重什么']?`看重：${{esc(r['买之前看重什么'])}}<br>`:''}}${{r['好评的事']?`<span class="g">好：${{esc(r['好评的事'])}}</span><br>`:''}}${{r['差评的事']?`<span class="b">差：${{esc(r['差评的事'])}}</span>`:''}}</div></div>`).join('');}}
render();
</script></body></html>'''
    open(path, "w", encoding="utf-8").write(page)


def build_xlsx(d, m, path, product, ins):
    wb = Workbook()
    ws = wb.active
    ws.title = "打标明细"
    ws.append(list(d.columns))
    for c in ws[1]:
        c.font = F(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for _, r in d.iterrows():
        ws.append(list(r.values))
    ws.freeze_panes = "B2"
    s = wb.create_sheet("分流与排名")
    s["A1"] = f"{product}｜Reddit 讨论 · 分流与排名"
    s["A1"].font = F(bold=True, size=13, color="1F3864")
    r = 3
    s.cell(r, 1, "到哪一步 × 说谁的货").font = F(bold=True)
    r += 1
    for j, w in enumerate(["到哪一步"] + m["whose_order"], 1):
        s.cell(r, j, w).font = F(bold=True)
    for st in m["stages"]:
        r += 1
        s.cell(r, 1, st)
        for j, w in enumerate(m["whose_order"], 2):
            s.cell(r, j, m["matrix"][st][w])
    for title, key in (("想解决什么", "need"), ("买之前看重什么（还没买的人）", "care"), ("现在拿什么凑合", "workaround"),
                       ("放弃的人：为什么放弃", "quit_why")):
        r += 2
        s.cell(r, 1, title).font = F(bold=True)
        r += 1
        for j, h in enumerate(["项", "条数", "点赞合计"], 1):
            s.cell(r, j, h).font = F(bold=True)
        for x in m[key]:
            r += 1
            s.cell(r, 1, x["tag"]); s.cell(r, 2, x["n"]); s.cell(r, 3, x["votes"])
    for title, key in (("买了的人 · 说我的货", "mine"), ("买了的人 · 说对手的货和整个品类", "rival")):
        r += 2
        s.cell(r, 1, title).font = F(bold=True)
        r += 1
        for j, h in enumerate(["细项", "好评", "差评"], 1):
            s.cell(r, j, h).font = F(bold=True)
        for x in m[key]:
            r += 1
            s.cell(r, 1, x["tag"]); s.cell(r, 2, x["好"]); s.cell(r, 3, x["差"])
    s.column_dimensions["A"].width = 34
    p = wb.create_sheet("客户的说法")
    p.append(["说法", "出现次数"])
    for x in m["phrases"]:
        p.append([x["p"], x["n"]])
    p.column_dimensions["A"].width = 44
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tagged"); ap.add_argument("insights")
    ap.add_argument("--config", required=True); ap.add_argument("--prep", default=None)
    ap.add_argument("--sku", default=None); ap.add_argument("--product", default=None)
    ap.add_argument("--outdir", default=None); ap.add_argument("--sample", action="store_true")
    a = ap.parse_args()
    cfg = load_config(a.config)
    d = pd.read_csv(a.tagged, dtype=str).fillna("")
    sku = json.load(open(a.sku, encoding="utf-8")) if a.sku else None
    validate(d, cfg, sku)
    prep = json.load(open(a.prep, encoding="utf-8")) if a.prep else {}
    ins = json.load(open(a.insights, encoding="utf-8"))
    product = a.product or ins.get("product") or cfg["品类"]
    outdir = a.outdir or ("/mnt/user-data/outputs" if os.path.isdir("/mnt/user-data/outputs") else "outputs")
    os.makedirs(outdir, exist_ok=True)
    seqs = set(d["序号"].astype(str))
    for t in ins.get("todo", []):
        if t.get("refs"):
            bad = [r for r in t["refs"] if str(r) not in seqs]
            if bad:
                sys.exit(f"[行动清单] 「{t.get('problem')}」依据的原话编号不存在：{bad}")
            t["n"] = len(t["refs"])
    m = compute(d, prep, cfg)
    if m["n"] < 30:
        print("[警告] 不到 30 条。只通读，不要在报告里出任何比例。", file=sys.stderr)
    safe = product.replace("/", "-")
    html_ = os.path.join(outdir, f"{safe}_Reddit讨论洞察报告.html")
    xlsx = os.path.join(outdir, f"{safe}_Reddit讨论分析.xlsx")
    build_html(d, m, ins, html_, product, cfg, a.sample)
    build_xlsx(d, m, xlsx, product, ins)
    dump({k: v for k, v in m.items() if k != "prep"}, os.path.join(outdir, "stats.json"))
    print(dump({"html": os.path.abspath(html_), "xlsx": os.path.abspath(xlsx), "进入分析": m["n"], "广告跑题": m["noise"],
                "到哪一步": m["stage"], "说谁的货": m["whose"]}))


if __name__ == "__main__":
    main()
