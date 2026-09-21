#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
退货线 · 第 3 步 算数，出报告

进来:
  tagged.csv     逐条打标后的表（列名见 references/02-打标表的格子.md）
  insights.json  叙述部分，由 AI 写（格式见 references/04-洞察怎么写.md）
  --prep         prep_return.py 出的数据体检
  --config       品类配置
  --sku          这一款的特点（人填的那张表），里面可以声明这一款特有的标签
出去:
  <产品>_退货分析.xlsx    统计全部用公式连着明细，业务自己改一个标签，数字跟着变
  <产品>_退货分诊报告.html  可以直接打印成 PDF

分工：脚本不编故事，AI 不算数字。
"""
import argparse, json, os, sys
from voc_common import (pd, load_config, multi, split_cell, base_tag, tag_dir, bands_of,
                        symmetry, all_tags, dir_defs, SEP, DIR_SEP, dump)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as GL

NEED = ["序号", "平台原因码", "原话", "哪件事_大类", "哪件事_细项", "谁在说", "什么情况下",
        "过程感受", "该谁改", "多严重", "能否挽回"]
HIDE_IN_HTML = ["记录编号", "订单号"]          # 网页里不带能对到人的列

F = lambda **k: Font(name="Arial", **k)
HDR = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E2F3")
_t = Side(style="thin", color="BFBFBF")
BD = Border(left=_t, right=_t, top=_t, bottom=_t)


# ------------------------------------------------------------ 校验
def validate(d, cfg, sku):
    err = []
    miss = [c for c in NEED if c not in d.columns]
    if miss:
        sys.exit(f"[打标表缺格子] {miss}\n格子的标准叫法见 references/02-打标表的格子.md")
    attrs = [a["名称"] for a in cfg["该谁改"]]
    bad = sorted(set(d["该谁改"]) - set(attrs))
    if bad:
        err.append(f"「该谁改」只能填 {attrs}，出现了 {bad}")
    bad = sorted(set(d["多严重"]) - set(cfg["多严重"]))
    if bad:
        err.append(f"「多严重」只能填 {cfg['多严重']}，出现了 {bad}")
    bad = sorted(set(d["能否挽回"]) - {"是", "否"})
    if bad:
        err.append(f"「能否挽回」只能填 是 / 否，出现了 {bad}")
    bad = sorted(set(d["哪件事_大类"]) - set(cfg["哪件事"]["大类顺序"]))
    if bad:
        err.append(f"「哪件事_大类」出现了标签表里没有的：{bad}")

    ok, dirs = all_tags(cfg, (sku or {}).get("这一款特有的标签")), dir_defs(cfg)
    unknown, bad_dir = set(), set()
    for v in d["哪件事_细项"]:
        for t in split_cell(v):
            b, dr = base_tag(t), tag_dir(t)
            if b not in ok:
                unknown.add(b)
            elif dr and (b not in dirs or dr not in dirs[b]):
                bad_dir.add(t)
    if unknown:
        err.append("这些细项不在标签表里（标签表定了就不改；确实是新东西，先记成下一版的提案，"
                   "这一轮放进兜底的那一项）：\n    " + "、".join(sorted(unknown)))
    if bad_dir:
        err.append("这些细项的方向写法不对（只有标签表里声明了两头的细项才能带方向）：\n    "
                   + "、".join(sorted(bad_dir)))
    for v in list(d["谁在说"]) + list(d["什么情况下"]) + list(d["过程感受"]):
        for t in split_cell(v):
            if t in ("无", "NA", "N/A", "none"):
                err.append("没内容的格子要留空，不要写「无」「NA」——会被当成一个标签统计进去")
                break
    if err:
        sys.exit("[打标表没过校验]\n- " + "\n- ".join(dict.fromkeys(err)))


# ------------------------------------------------------------ 算数
def count_dir(d, name, direction):
    target = f"{name}{DIR_SEP}{direction}"
    return int(d["哪件事_细项"].map(lambda v: target in split_cell(v)).sum())


def compute(d, prep, cfg):
    n = len(d)
    bands = bands_of(cfg)
    full = multi(d["哪件事_细项"])
    base = full.groupby(full.index.map(base_tag)).sum().sort_values(ascending=False)
    gate = max(3, round(n * 0.03))
    dirs = dir_defs(cfg)
    overall_names = [a["名称"] for a in cfg.get("两头属性", [])]

    # 整体那一层（原话层）
    sym_overall = []
    for att in cfg.get("两头属性", []):
        nm, a, b = att["名称"], att["一头"]["叫法"], att["另一头"]["叫法"]
        s = symmetry(count_dir(d, nm, a), count_dir(d, nm, b), a, b, bands)
        s["属性"] = nm
        sym_overall.append(s)

    # 逐项那一层：同一把尺子到下一层再量
    item_min = int(cfg.get("逐项最少条数", 8))
    sym_items = []
    for nm, (a, b) in dirs.items():
        if nm in overall_names:
            continue
        ca, cb = count_dir(d, nm, a), count_dir(d, nm, b)
        tot = ca + cb
        if not tot:
            continue
        pa = ca / tot * 100
        skew = abs(pa - (100 - pa))
        if tot < item_min:
            verdict = "样本少"
        elif max(pa, 100 - pa) >= 90:
            verdict = "一边倒"
        elif skew <= 30:
            verdict = "两头都有"
        else:
            verdict = "往一头偏"
        sym_items.append({"细项": nm, "一头": a, "另一头": b, "一头条数": ca, "另一头条数": cb,
                          "合计": tot, "一头占比": round(pa, 1), "另一头占比": round(100 - pa, 1),
                          "差距": round(skew, 1), "判定": verdict})
    sym_items.sort(key=lambda x: -x["合计"])

    # 原因码方向 vs 原话方向
    rvv = {}
    for att in cfg.get("两头属性", []):
        for side, other in (("一头", "另一头"), ("另一头", "一头")):
            for code in att[side].get("原因码", []):
                g = d[d["平台原因码"] == code]
                if not len(g):
                    continue
                same = count_dir(g, att["名称"], att[side]["叫法"])
                opp = count_dir(g, att["名称"], att[other]["叫法"])
                rvv[code] = {"条数": len(g), "方向一致": same, "方向相反": opp,
                             "未表方向": len(g) - same - opp, "一致率": round(same / len(g), 3)}

    # 量小但严重
    top_sev = cfg["多严重"][0]
    sev_total = int((d["多严重"] == top_sev).sum())
    sev_by_l1 = [{"l1": l1, "n": int((d["哪件事_大类"] == l1).sum()),
                  "hot": int(((d["哪件事_大类"] == l1) & (d["多严重"] == top_sev)).sum())}
                 for l1 in cfg["哪件事"]["大类顺序"]]
    for x in sev_by_l1:
        x["占全部严重"] = round(x["hot"] / sev_total, 3) if sev_total else 0

    # 被平台字段盖住的：每个进了排名的细项，平台把它记成了什么
    # 一个规格还是整条线：这个细项在各规格里的占比，跟大盘比
    tag_by_reason, tag_by_spec = {}, {}
    spec_all = d["规格"].value_counts(normalize=True).to_dict() if "规格" in d.columns else {}
    for t in base[base >= gate].index:
        hit = d[d["哪件事_细项"].map(lambda v: t in [base_tag(x) for x in split_cell(v)])]
        tag_by_reason[t] = hit["平台原因码"].value_counts().to_dict()
        if spec_all:
            sp = hit["规格"].value_counts(normalize=True).to_dict()
            tag_by_spec[t] = {k: {"这一项里占": round(sp.get(k, 0), 3), "大盘里占": round(v, 3)}
                              for k, v in spec_all.items()}

    # 大筐拆开：每个大筐码里，客户实际在说哪几类事
    basket = {c: d.loc[d["平台原因码"] == c, "哪件事_大类"].value_counts().to_dict()
              for c in cfg["原因码"].get("大筐", []) if (d["平台原因码"] == c).any()}

    return {
        "n": n, "gate": gate,
        "l1": d["哪件事_大类"].value_counts().to_dict(),
        "tags": base.to_dict(), "tags_full": full.to_dict(),
        "tags_in": base[base >= gate].to_dict(),
        "attr": d["该谁改"].value_counts().to_dict(),
        "sev": d["多严重"].value_counts().to_dict(),
        "sev_total": sev_total, "sev_by_l1": sev_by_l1,
        "save": int((d["能否挽回"].astype(str).str.strip() == "是").sum()),
        "reason": d["平台原因码"].value_counts().to_dict(),
        "spec": d["规格"].value_counts().to_dict() if "规格" in d.columns else {},
        "month": d["月份"].value_counts().sort_index().to_dict() if "月份" in d.columns else {},
        "user": multi(d["谁在说"]).to_dict(),
        "scene": multi(d["什么情况下"]).to_dict(),
        "exp": multi(d["过程感受"]).to_dict(),
        "links": [{"r": r, "a": a, "v": int(v)}
                  for (r, a), v in d.groupby(["平台原因码", "该谁改"]).size().items()],
        "sym_overall": sym_overall, "sym_items": sym_items, "rvv": rvv,
        "tag_by_reason": tag_by_reason, "tag_by_spec": tag_by_spec, "basket": basket,
        "prep": prep or {},
    }


# ------------------------------------------------------------ Excel
def head(ws, row, labels):
    for i, v in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = F(bold=True, color="FFFFFF", size=10)
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BD


def body(ws, r1, r2, c2, wrap=False):
    for r in range(r1, r2 + 1):
        for c in range(1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = F(size=10)
            cell.border = BD
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)


def title(ws, text):
    ws["A1"] = text
    ws["A1"].font = F(bold=True, size=12, color="1F3864")


def xl(s):
    """放进 Excel 公式里的字符串：双引号翻倍；SEARCH 的通配符转义。"""
    return str(s).replace('"', '""').replace("~", "~~").replace("?", "~?").replace("*", "~*")


def build_xlsx(d, m, path, product, ins, cfg):
    wb = Workbook()
    n, p = m["n"], m["prep"]
    attrs = cfg["该谁改"]

    ws = wb.active
    ws.title = "00_说明与口径"
    ws["A1"] = f"{product}｜退货分析"
    ws["A1"].font = F(bold=True, size=14, color="1F3864")
    rows = [("", "", ""), ("一、清洗漏斗", "", ""), ("步骤", "规则", "剩下")]
    if p:
        sc, lt, vo = p.get("本次范围", {}), p.get("留剔", {}), p.get("原话", {})
        rows += [("原始行数", "导出文件的全部行", p.get("原始行数")),
                 ("只留退货", "换货用的是另一套原因码，必须先筛掉", p.get("退货行数")),
                 ("限定范围", f"产品={sc.get('产品')} · 站点={sc.get('站点')}", sc.get("行数")),
                 ("剔掉与产品无关的", f"物流 / 客户改主意 / 没给原因，共 {lt.get('剔掉的码有几类')} 类", lt.get("留")),
                 ("留下且有原话", f"没写字的、只有平台套话的（{vo.get('只有平台套话被剔掉')} 条）都去掉", n)]
        k = p.get("口径", {})
        rows += [("", "", ""), ("二、口径", "", ""),
                 ("退货明细行", "本次范围内的退货行数", k.get("退货明细行")),
                 ("退货订单", "同一订单只算一次", k.get("退货订单")),
                 ("写了原因的占比", "写了字的 ÷ 明细行", vo.get("覆盖率"))]
        for key, val in (p.get("提醒") or {}).items():
            rows += [(f"⚠ {key}", val if isinstance(val, str) else json.dumps(val, ensure_ascii=False), "")]
    rows += [("", "", ""), ("三、「该谁改」怎么判", "", "")]
    rows += [(a["名称"], a["判法"] + "  →  " + a["方向"], a["部门"]) for a in attrs]
    rows += [("", "", ""), ("四、这批数据答不了什么", "", "")]
    rows += [("", c, "") for c in (ins.get("caveats") or [])]
    r = 2
    for row in rows:
        for i, v in enumerate(row, 1):
            ws.cell(row=r, column=i, value=v if v != "" else None)
        if isinstance(row[0], str) and row[0][:2] in ("一、", "二、", "三、", "四、"):
            ws.cell(row=r, column=1).font = F(bold=True, size=11, color="1F3864")
        r += 1
    for rr in range(2, r):
        for c in range(1, 4):
            cell = ws.cell(row=rr, column=c)
            if not cell.font.bold:
                cell.font = F(size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for c_, w in zip("ABC", (24, 76, 18)):
        ws.column_dimensions[c_].width = w

    # 01 打标明细
    ws = wb.create_sheet("01_打标明细")
    cols = [c for c in ["序号", "记录编号", "订单号", "站点", "规格", "月份", "平台原因码", "原话",
                        "哪件事_大类", "哪件事_细项", "谁在说", "什么情况下", "过程感受",
                        "该谁改", "多严重", "能否挽回"] if c in d.columns]
    ws.append(cols)
    head(ws, 1, cols)
    for _, row in d[cols].iterrows():
        ws.append(list(row.values))
    last = n + 1
    body(ws, 2, last, len(cols), wrap=True)
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[GL(i)].width = {"原话": 58, "哪件事_细项": 36, "谁在说": 24, "什么情况下": 20,
                                             "过程感受": 20, "平台原因码": 24, "记录编号": 18,
                                             "订单号": 20, "规格": 16}.get(c, 12)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{GL(len(cols))}{last}"

    DET = "'01_打标明细'"
    col_of = {c: GL(i) for i, c in enumerate(cols, 1)}
    rng = lambda c: f"{DET}!${col_of[c]}$2:${col_of[c]}${last}"
    cnt = lambda tag, c: f'=SUMPRODUCT(--ISNUMBER(SEARCH("{xl(tag)}",{rng(c)})))'

    # 02 标签统计
    ws = wb.create_sheet("02_标签统计")
    title(ws, f"客户在说什么（一条原话可以说到好几件事，合计 > {n}）｜进排名的门槛 ≥{m['gate']}")
    head(ws, 3, ["大类", "细项", "提到的条数", f"占 {n} 条"])
    r = 4
    for l1 in [x for x in cfg["哪件事"]["大类顺序"] if x in m["l1"]]:
        ws.cell(row=r, column=1, value=l1)
        ws.cell(row=r, column=2, value="— 小计")
        ws.cell(row=r, column=3, value=f'=COUNTIF({rng("哪件事_大类")},"{xl(l1)}")')
        ws.cell(row=r, column=4, value=f"=C{r}/{n}")
        for c in range(1, 5):
            ws.cell(row=r, column=c).fill = SUB
            ws.cell(row=r, column=c).font = F(bold=True, size=10)
        r += 1
        here = d.loc[d["哪件事_大类"] == l1, "哪件事_细项"]
        subs = [t for t in m["tags"] if here.map(lambda v: t in [base_tag(x) for x in split_cell(v)]).any()]
        for t in sorted(subs, key=lambda x: -m["tags"][x]):
            ws.cell(row=r, column=2, value=t)
            ws.cell(row=r, column=3, value=cnt(t, "哪件事_细项"))
            ws.cell(row=r, column=4, value=f"=C{r}/{n}")
            r += 1
    for grp, key in (("谁在说", "user"), ("什么情况下", "scene"), ("过程感受", "exp")):
        r += 1
        ws.cell(row=r, column=1, value=grp).font = F(bold=True, size=11, color="1F3864")
        r += 1
        head(ws, r, ["类型", "标签", "提到的条数", f"占 {n} 条"])
        r += 1
        for t, _v in sorted(m[key].items(), key=lambda kv: -kv[1]):
            ws.cell(row=r, column=1, value=grp)
            ws.cell(row=r, column=2, value=t)
            ws.cell(row=r, column=3, value=cnt(t, grp))
            ws.cell(row=r, column=4, value=f"=C{r}/{n}")
            r += 1
    body(ws, 4, r - 1, 4)
    for rr in range(4, r):
        ws.cell(row=rr, column=4).number_format = "0.0%"
        ws.cell(row=rr, column=3).alignment = Alignment(horizontal="center")
    for c_, w in zip("ABCD", (20, 32, 14, 12)):
        ws.column_dimensions[c_].width = w

    # 03 该谁改与严重度
    ws = wb.create_sheet("03_该谁改与严重度")
    title(ws, "该谁改 —— 每条原话只判一个")
    head(ws, 3, ["该谁改", "条数", "占比", "哪个部门", "往哪个方向动"])
    r, total_row = 4, 4 + len(attrs)
    for a in attrs:
        ws.cell(row=r, column=1, value=a["名称"])
        ws.cell(row=r, column=2, value=f'=COUNTIF({rng("该谁改")},"{xl(a["名称"])}")')
        ws.cell(row=r, column=3, value=f"=B{r}/$B${total_row}")
        ws.cell(row=r, column=4, value=a["部门"])
        ws.cell(row=r, column=5, value=a["方向"])
        r += 1
    ws.cell(row=r, column=1, value="合计")
    ws.cell(row=r, column=2, value=f"=SUM(B4:B{r-1})")
    for c in range(1, 6):
        ws.cell(row=r, column=c).fill = SUB
        ws.cell(row=r, column=c).font = F(bold=True, size=10)
    body(ws, 4, r, 5, wrap=True)
    for rr in range(4, r + 1):
        ws.cell(row=rr, column=3).number_format = "0.0%"
    for label, colname, values in (("多严重", "多严重", cfg["多严重"]), ("能否挽回", "能否挽回", ["是", "否"])):
        r += 2
        ws.cell(row=r, column=1, value=label).font = F(bold=True, size=11, color="1F3864")
        r += 1
        head(ws, r, [label, "条数", "占比"])
        r += 1
        s0 = r
        for v in values:
            ws.cell(row=r, column=1, value=v)
            ws.cell(row=r, column=2, value=f'=COUNTIF({rng(colname)},"{xl(v)}")')
            ws.cell(row=r, column=3, value=f"=B{r}/{n}")
            r += 1
        body(ws, s0, r - 1, 3)
        for rr in range(s0, r):
            ws.cell(row=rr, column=3).number_format = "0.0%"
    r += 2
    ws.cell(row=r, column=1, value="量小不等于不要紧：每一大类里有几条是「严重」").font = F(bold=True, size=11, color="1F3864")
    r += 1
    head(ws, r, ["大类", "条数", "其中严重", "占全部严重"])
    r += 1
    s0 = r
    top = cfg["多严重"][0]
    for x in m["sev_by_l1"]:
        ws.cell(row=r, column=1, value=x["l1"])
        ws.cell(row=r, column=2, value=f'=COUNTIF({rng("哪件事_大类")},"{xl(x["l1"])}")')
        ws.cell(row=r, column=3, value=f'=COUNTIFS({rng("哪件事_大类")},"{xl(x["l1"])}",{rng("多严重")},"{xl(top)}")')
        ws.cell(row=r, column=4, value=f'=IF(COUNTIF({rng("多严重")},"{xl(top)}")=0,0,C{r}/COUNTIF({rng("多严重")},"{xl(top)}"))')
        r += 1
    body(ws, s0, r - 1, 4)
    for rr in range(s0, r):
        ws.cell(row=rr, column=4).number_format = "0.0%"
    for c_, w in zip("ABCDE", (24, 12, 12, 22, 54)):
        ws.column_dimensions[c_].width = w

    # 04 哪件事 × 规格
    if "规格" in d.columns:
        ws = wb.create_sheet("04_哪件事×规格")
        title(ws, "细项 × 规格（数字=条数）｜用来判断：是某个规格的事，还是整条线的事。分组小的格子只看方向")
        pref = [s for s in cfg.get("规格排序", []) if s in m["spec"]]
        specs = pref + [s for s in m["spec"] if s not in pref]
        head(ws, 3, ["细项"] + specs + ["合计"])
        r = 4
        for t in sorted(m["tags"], key=lambda x: -m["tags"][x]):
            ws.cell(row=r, column=1, value=t)
            for i, s in enumerate(specs):
                ws.cell(row=r, column=2 + i,
                        value=f'=SUMPRODUCT(--ISNUMBER(SEARCH("{xl(t)}",{rng("哪件事_细项")})),'
                              f'--({rng("规格")}="{xl(s)}"))')
            ws.cell(row=r, column=2 + len(specs), value=f"=SUM(B{r}:{GL(1+len(specs))}{r})")
            r += 1
        ws.cell(row=r, column=1, value="这个规格的原话总数（大盘）")
        for i, s in enumerate(specs):
            ws.cell(row=r, column=2 + i, value=f'=COUNTIF({rng("规格")},"{xl(s)}")')
        for c in range(1, 3 + len(specs)):
            ws.cell(row=r, column=c).fill = SUB
            ws.cell(row=r, column=c).font = F(bold=True, size=10)
        body(ws, 4, r, 2 + len(specs))
        ws.column_dimensions["A"].width = 30
        for i in range(len(specs) + 1):
            ws.column_dimensions[GL(2 + i)].width = 15
        ws.freeze_panes = "B4"

    # 05 该谁改 × 原因码
    ws = wb.create_sheet("05_该谁改×原因码")
    title(ws, "该谁改 × 平台原因码（看平台记的原因，和读出来的真实原因差多远）")
    reasons = sorted(m["reason"], key=lambda x: -m["reason"][x])
    head(ws, 3, ["该谁改"] + [f"{p.get('原因码中文', {}).get(c, c)}\n{c}" for c in reasons] + ["合计"])
    r = 4
    for a in attrs:
        ws.cell(row=r, column=1, value=a["名称"])
        for i, rs in enumerate(reasons):
            ws.cell(row=r, column=2 + i, value=f'=SUMPRODUCT(--({rng("该谁改")}="{xl(a["名称"])}"),'
                                               f'--({rng("平台原因码")}="{xl(rs)}"))')
        ws.cell(row=r, column=2 + len(reasons), value=f"=SUM(B{r}:{GL(1+len(reasons))}{r})")
        r += 1
    ws.cell(row=r, column=1, value="合计")
    for i, rs in enumerate(reasons):
        ws.cell(row=r, column=2 + i, value=f'=COUNTIF({rng("平台原因码")},"{xl(rs)}")')
    ws.cell(row=r, column=2 + len(reasons), value=f"=SUM(B{r}:{GL(1+len(reasons))}{r})")
    for c in range(1, 3 + len(reasons)):
        ws.cell(row=r, column=c).fill = SUB
        ws.cell(row=r, column=c).font = F(bold=True, size=10)
    body(ws, 4, r, 2 + len(reasons))
    ws.column_dimensions["A"].width = 22
    for i in range(len(reasons) + 1):
        ws.column_dimensions[GL(2 + i)].width = 18
        ws.cell(row=3, column=2 + i).alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[3].height = 42
    ws.freeze_panes = "B4"

    # 06 原因码大盘与两头对称
    ws = wb.create_sheet("06_大盘与两头对称")
    title(ws, "平台原因码的大盘（客户在下拉框里随手选的，只作参考，不拿来判该谁改）")
    rd, rs_ = p.get("原因码分布", {}), p.get("原因码占比", {})
    keepset = set(p.get("留下的码", []))
    r = 3
    if rd:
        head(ws, r, ["原因码", "中文", "条数", "占本次范围", "处理"])
        r += 1
        r0 = r
        for code, v in sorted(rd.items(), key=lambda kv: -kv[1]):
            for i, val in enumerate([code, p.get("原因码中文", {}).get(code, ""), v, rs_.get(code, ""),
                                     "留" if code in keepset else "剔"], 1):
                ws.cell(row=r, column=i, value=val)
            r += 1
        body(ws, r0, r - 1, 5)
        for rr in range(r0, r):
            ws.cell(row=rr, column=4).number_format = "0.0%"
    r += 2
    ws.cell(row=r, column=1, value="两头是否对称（整体一层，再到每一项）").font = F(bold=True, size=11, color="1F3864")
    r += 1
    head(ws, r, ["哪一层 / 哪一项", "一头", "条数", "另一头", "条数", "差距(点)", "判定"])
    r += 1
    r0 = r
    for s in (p.get("两头对称_原因码层") or []):
        if s.get("合计"):
            for i, val in enumerate([f"{s['属性']}·原因码层", s["一头"], s["一头条数"], s["另一头"],
                                     s["另一头条数"], s["差距"], s["判定"]], 1):
                ws.cell(row=r, column=i, value=val)
            r += 1
    for s in m["sym_overall"]:
        if s.get("合计"):
            for i, val in enumerate([f"{s['属性']}·原话层", s["一头"], s["一头条数"], s["另一头"],
                                     s["另一头条数"], s["差距"], s["判定"]], 1):
                ws.cell(row=r, column=i, value=val)
            r += 1
    for s in m["sym_items"]:
        for i, val in enumerate([f"逐项·{s['细项']}", s["一头"], s["一头条数"], s["另一头"],
                                 s["另一头条数"], s["差距"], s["判定"]], 1):
            ws.cell(row=r, column=i, value=val)
        r += 1
    if r > r0:
        body(ws, r0, r - 1, 7)
    for c_, w in zip("ABCDEFG", (30, 16, 10, 16, 10, 12, 14)):
        ws.column_dimensions[c_].width = w
    wb.save(path)


# ------------------------------------------------------------ HTML
TPL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "return_report_template.html")


def build_html(d, m, ins, path, product, cfg, is_sample):
    tpl = open(TPL, encoding="utf-8").read()
    recs = d.drop(columns=[c for c in HIDE_IN_HTML if c in d.columns]).fillna("").to_dict("records")
    payload = {
        "product": product, "is_sample": bool(is_sample), "stats": m, "records": recs, "insights": ins,
        "cfg": {"attrs": [{"name": a["名称"], "color": a["颜色"], "dept": a["部门"], "dir": a["方向"]}
                          for a in cfg["该谁改"]],
                "texts": cfg.get("图上文字", {}), "bands": cfg.get("对称档位", []),
                "item_min_n": cfg.get("逐项最少条数", 8)},
    }
    out = tpl.replace("__PAYLOAD__", json.dumps(payload, ensure_ascii=False, default=str).replace("</", "<\\/"))
    open(path, "w", encoding="utf-8").write(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("tagged")
    ap.add_argument("insights")
    ap.add_argument("--config", required=True)
    ap.add_argument("--prep", default=None)
    ap.add_argument("--sku", default=None, help="这一款的特点（人填的表，JSON）")
    ap.add_argument("--outdir", default=None)
    ap.add_argument("--sample", action="store_true", help="在报告上标明「示例数据」")
    a = ap.parse_args()

    cfg = load_config(a.config)
    d = pd.read_csv(a.tagged, dtype=str).fillna("")
    sku = json.load(open(a.sku, encoding="utf-8")) if a.sku else None
    if not sku:
        print("[提醒] 没给「这一款的特点」。卖点同时也是漏点这类结论，没有它看不出来。", file=sys.stderr)
    validate(d, cfg, sku)

    prep = json.load(open(a.prep, encoding="utf-8")) if a.prep else {}
    ins = json.load(open(a.insights, encoding="utf-8"))
    product = ins.get("product") or (prep.get("本次范围", {}) or {}).get("产品") or "未命名"
    outdir = a.outdir or ("/mnt/user-data/outputs" if os.path.isdir("/mnt/user-data/outputs") else "outputs")
    os.makedirs(outdir, exist_ok=True)

    # 行动依据多少条，由脚本按「依据的原话编号」去数，不让 AI 自己填数字
    seqs = set(d["序号"].astype(str))
    for t in ins.get("todo", []):
        if t.get("refs"):
            bad = [r for r in t["refs"] if str(r) not in seqs]
            if bad:
                sys.exit(f"[行动清单] 「{t.get('problem') or t.get('id')}」依据的原话编号不存在：{bad}")
            t["n"] = len(t["refs"])
        elif t.get("n") is None:
            print(f"[提醒] 行动「{t.get('problem') or t.get('id')}」没写依据的原话编号（refs），没法点回原话。",
                  file=sys.stderr)

    m = compute(d, prep, cfg)
    if m["n"] < 30:
        print("[警告] 有原话的不到 30 条。只通读，不要在报告里出任何比例。", file=sys.stderr)

    safe = product.replace("/", "-")
    xlsx = os.path.join(outdir, f"{safe}_退货分析.xlsx")
    html_ = os.path.join(outdir, f"{safe}_退货分诊报告.html")
    build_xlsx(d, m, xlsx, product, ins, cfg)
    build_html(d, m, ins, html_, product, cfg, a.sample)
    dump({k: m[k] for k in m if k not in ("prep",)}, os.path.join(outdir, "stats.json"))

    print(dump({"xlsx": os.path.abspath(xlsx), "html": os.path.abspath(html_),
                "有原话的": m["n"], "进排名的门槛": m["gate"], "该谁改": m["attr"],
                "严重": m["sev_total"], "能挽回": m["save"],
                "整体两头": [(s["属性"], s.get("判定"), s.get("差距")) for s in m["sym_overall"]],
                "逐项两头": [(s["细项"], s["判定"], f"{s['一头条数']}:{s['另一头条数']}") for s in m["sym_items"]]}))
    print("\n下一步：python recalc.py " + os.path.abspath(xlsx))


if __name__ == "__main__":
    main()
